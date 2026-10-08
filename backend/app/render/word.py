"""DOCX export (DESIGN.md §11.6, EXP-2): each template's résumé and cover letter as an editable Word document.

The builders use the same views as the HTML templates (html.py), so the content, its order and its formatting
(dates, grouped roles, links) match the PDF. The parity target is "the same design language, slightly
simplified": Meridian and Ledger are single-column with tab-stop dates, Atlas a borderless two-column table with
a shaded sidebar, and Vanguard a shaded header band with section labels in a gutter.
"""

from __future__ import annotations

from typing import Any

from docx.enum.text import WD_BREAK
from docx.shared import Emu, Inches, Pt, RGBColor

from app.render.html import letter_view, resume_view
from app.render.theme import Theme
from app.render.word_kit import MUTED, Kit, mix

SKILL_LABEL = Inches(1.55)
GUTTER, GAP = Inches(1.0), Pt(14)


# ---- résumé content -------------------------------------------------------------------------------------


def _entry_lines(kit: Kit, container: Any, e: dict[str, Any], style: str = "Body") -> None:
    paragraph = container.add_paragraph(style=style)
    paragraph.paragraph_format.keep_together = True
    title = paragraph.add_run(e["title"])
    title.bold = True
    if e.get("meta"):
        paragraph.add_run().add_break(WD_BREAK.LINE)
        paragraph.add_run(e["meta"], style=kit.doc.styles["Muted"])
    if e.get("summary"):
        kit.para(container, "Body", e["summary"])
    kit.bullets(container, e.get("bullets") or [])


def _org(kit: Kit, container: Any, group: dict[str, Any]) -> None:
    paragraph = container.add_paragraph(style="Org")
    paragraph.add_run(group["company"])
    if group.get("location"):
        paragraph.add_run(f" · {group['location']}", style=kit.doc.styles["Muted"]).bold = False


def _section(kit: Kit, container: Any, section: dict[str, Any], width: Emu | int, *, heading: bool = True) -> None:
    """One section's content (and, unless the layout puts it elsewhere, its heading)."""
    kind = section["kind"]
    if kind == "earlier-experience":
        paragraph = container.add_paragraph(style="Body")
        paragraph.paragraph_format.space_before = Pt(6)
        if heading:
            paragraph.add_run(f"{section['title']}: ").bold = True
        paragraph.add_run("; ".join(section["entries"]), style=kit.doc.styles["Muted"])
        return
    if heading:
        kit.para(container, "Section Heading", section["title"], inline=False)
    ledger = kit.template == "ledger"
    if kind == "experience":
        for n, group in enumerate(section["groups"]):
            if ledger:
                first = kit.row(container, "Org", group["company"], group.get("location") or "", width)
                for role in group["roles"]:
                    title = kit.row(container, "Body", role["position"], role["dates"], width)
                    title.runs[0].italic = True
                    title.paragraph_format.keep_with_next = True
                    kit.bullets(container, role["bullets"])
            elif len(group["roles"]) == 1:
                role = group["roles"][0]
                first = kit.row(container, "Role Title", role["position"], role["dates"], width)
                _org(kit, container, group)
                kit.bullets(container, role["bullets"])
            else:
                _org(kit, container, group)
                first = container.paragraphs[-1]
                for role in group["roles"]:
                    kit.row(container, "Role Title", role["position"], role["dates"], width)
                    kit.bullets(container, role["bullets"])
            if n:
                first.paragraph_format.space_before = Pt(8 * kit.space)
    elif kind == "projects":
        for n, p in enumerate(section["entries"]):
            name = p["name"] + (f" · {p['aside']}" if p.get("aside") else "")
            row = kit.row(container, "Org" if ledger else "Role Title", name, p.get("right") or "", width)
            if n:
                row.paragraph_format.space_before = Pt(6 * kit.space)
            kit.bullets(container, p["bullets"])
    elif kind == "skills":
        for group in section["groups"]:
            paragraph = container.add_paragraph(style="Body")
            fmt = paragraph.paragraph_format
            fmt.left_indent, fmt.first_line_indent = SKILL_LABEL, -SKILL_LABEL
            fmt.tab_stops.add_tab_stop(SKILL_LABEL)
            paragraph.add_run(group["name"]).bold = True
            paragraph.add_run("\t" + group["text"])
    elif kind == "languages":
        kit.para(container, "Body", section["text"], inline=False)
    else:
        for e in section["entries"]:
            _entry_lines(kit, container, e)


def _header(kit: Kit, container: Any, view: dict[str, Any]) -> None:
    kit.para(container, "Name", view["name"], inline=False)
    if view["headline"]:
        kit.para(container, "Headline", view["headline"], inline=False)
    kit.contact(container, view["contact"])
    if kit.template == "ledger":  # the double rule under the header
        last = container.paragraphs[-1]
        kit.border(last._p.get_or_add_pPr(), "bottom", kit.accent, size=6, space=6, kind="double")  # noqa: SLF001
        last.paragraph_format.space_after = Pt(4)


def _single(kit: Kit, view: dict[str, Any]) -> None:
    """Meridian and Ledger: one column."""
    doc = kit.doc
    _header(kit, doc, view)
    if view["summary"]:
        kit.para(doc, "Section Heading", "Summary", inline=False)
        kit.para(doc, "Body", view["summary"])
    for section in view["sections"]:
        _section(kit, doc, section, kit.width)


def _atlas(kit: Kit, view: dict[str, Any]) -> None:
    """A borderless table: the shaded sidebar (contact, skills, credentials) beside the main column."""
    side_w = Inches(2.3)
    main_w = kit.width - side_w
    table = kit.table([side_w, main_w])
    side, main = table.rows[0].cells
    kit.cell_style(side, fill=kit.color("tint"), pad=(12, 12, 12, 12))
    kit.cell_style(main, pad=(0, 0, 0, 18))
    main_text = main_w - Pt(18)

    kit.first_paragraph(main, "Name").add_run(view["name"])
    if view["headline"]:
        kit.para(main, "Headline", view["headline"], inline=False)
    if view["summary"]:
        kit.para(main, "Body", view["summary"]).paragraph_format.space_before = Pt(8)
    for section in view["sections"]:
        if not section["side"]:
            _section(kit, main, section, main_text)

    heading = kit.first_paragraph(side, "Section Heading")
    heading.add_run("Contact")
    heading.paragraph_format.space_before = Pt(0)
    _sidebar_headings(kit, heading)
    kit.contact(side, view["contact"], style="Sidebar", stacked=True)
    for section in view["sections"]:
        if not section["side"]:
            continue
        _sidebar_headings(kit, kit.para(side, "Section Heading", section["title"], inline=False))
        if section["kind"] == "skills":
            for group in section["groups"]:
                label = kit.para(side, "Sidebar", group["name"], inline=False)
                label.runs[0].bold = True
                label.paragraph_format.space_before = Pt(4)
                kit.para(side, "Sidebar", ", ".join(group["items"]), inline=False)
        elif section["kind"] == "languages":
            kit.para(side, "Sidebar", section["text"], inline=False)
        else:
            for e in section["entries"]:
                _entry_lines(kit, side, {**e, "bullets": [], "summary": None}, style="Sidebar")


def _sidebar_headings(kit: Kit, paragraph: Any) -> None:
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor.from_string(kit.accent)


def _band(kit: Kit, view: dict[str, Any]) -> None:
    """Vanguard's header: a shaded accent band with the name, headline and contact details in white."""
    band = kit.table([kit.width])
    cell = band.rows[0].cells[0]
    kit.cell_style(cell, fill=kit.accent, pad=(16, 18, 14, 18))
    kit.first_paragraph(cell, "Name").add_run(view["name"])
    if view["headline"]:
        kit.para(cell, "Headline", view["headline"], inline=False)
    kit.contact(cell, view["contact"], color=mix(kit.theme.accent, 0.3))
    # Without a paragraph between them, Word and LibreOffice merge two tables into one
    spacer = kit.doc.add_paragraph()
    spacer.paragraph_format.space_before = spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(6)


def _vanguard(kit: Kit, view: dict[str, Any]) -> None:
    """The band, then one table row per section: the label in a gutter, the content beside it."""
    _band(kit, view)
    sections = ([{"kind": "summary", "title": "Profile"}] if view["summary"] else []) + view["sections"]
    content_w = kit.width - GUTTER
    table = kit.table([GUTTER, content_w], rows=len(sections))
    for n, (row, section) in enumerate(zip(table.rows, sections, strict=True)):
        label, body = row.cells
        rule = "E2E4F0" if n else None
        kit.cell_style(label, pad=(7 if n else 12, 6, 7, 0), top_border=rule)
        kit.cell_style(body, pad=(7 if n else 12, 0, 7, 8), top_border=rule)
        heading = kit.first_paragraph(label, "Section Heading")
        heading.add_run(section["title"])
        heading.paragraph_format.space_before = heading.paragraph_format.space_after = Pt(0)
        # "keep with next" on a whole row would keep every row with the next: the table would move as one block
        heading.paragraph_format.keep_with_next = False
        first = body.paragraphs[0]
        if section["kind"] == "summary":
            first.style = "Body"
            kit.runs(first, view["summary"])
            continue
        _section(kit, body, section, content_w - Pt(8), heading=False)
        if len(body.paragraphs) > 1 and not first.text:  # drop the cell's initial empty paragraph
            first._p.getparent().remove(first._p)  # noqa: SLF001


def build_resume_docx(resume: dict[str, Any], template: str, theme: Theme, metadata: dict[str, str]) -> bytes:
    kit = Kit(template, theme)
    view = resume_view(resume, theme, pair=False)
    {"meridian": _single, "ledger": _single, "atlas": _atlas, "vanguard": _vanguard}[template](kit, view)
    kit.footer(view["name"])
    return kit.save(metadata)


# ---- cover letter ----------------------------------------------------------------------------------------


def _letter_text(kit: Kit, container: Any, view: dict[str, Any], *, meta: bool = True) -> None:
    """Date, recipient and Re (unless the layout shows them itself), then the salutation, body and sign-off."""
    if meta:
        if view["date"]:
            date = kit.para(container, "Letter Body", view["date"], inline=False)
            date.paragraph_format.space_before = Pt(18)
            date.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
        if view["recipient"]:
            recipient = container.add_paragraph(style="Recipient")
            for i, line in enumerate(view["recipient"]):
                if i:
                    recipient.add_run().add_break(WD_BREAK.LINE)
                recipient.add_run(line).bold = i == 0
        if view["subject"]:
            kit.para(container, "Subject", f"Re: {view['subject']}", inline=False)
    if view["salutation"]:
        salutation = kit.para(container, "Letter Body", view["salutation"], inline=False)
        salutation.paragraph_format.space_before = Pt(10)
    for text in view["paragraphs"]:
        kit.para(container, "Letter Body", text)
    closing = kit.para(container, "Letter Body", view["closing"], inline=False)
    closing.paragraph_format.space_before = Pt(6)
    closing.paragraph_format.keep_with_next = True
    kit.para(container, "Signature", view["signature"], inline=False)


def build_letter_docx(letter: dict[str, Any], template: str, theme: Theme, metadata: dict[str, str]) -> bytes:
    kit = Kit(template, theme)
    view = letter_view(letter)
    if template == "atlas":
        side_w = Inches(2.3)
        table = kit.table([side_w, kit.width - side_w])
        side, main = table.rows[0].cells
        kit.cell_style(side, fill=kit.color("tint"), pad=(12, 12, 12, 12))
        kit.cell_style(main, pad=(0, 0, 0, 18))
        kit.first_paragraph(main, "Name").add_run(view["name"])
        if view["headline"]:
            kit.para(main, "Headline", view["headline"], inline=False)
        _letter_text(kit, main, view)
        heading = kit.first_paragraph(side, "Section Heading")
        heading.add_run("Contact")
        heading.paragraph_format.space_before = Pt(0)
        _sidebar_headings(kit, heading)
        kit.contact(side, view["contact"], style="Sidebar", stacked=True)
        if view["highlights"]:
            title = f"Highlights for {view['company']}" if view["company"] else "Highlights"
            _sidebar_headings(kit, kit.para(side, "Section Heading", title, inline=False))
            for h in view["highlights"]:
                value = kit.para(side, "Sidebar", h["value"], inline=False)
                value.paragraph_format.space_before = Pt(6)
                value.runs[0].bold, value.runs[0].font.size = True, Pt(16)
                value.runs[0].font.color.rgb = RGBColor.from_string(kit.accent)
                kit.para(side, "Sidebar", h["label"], inline=False)
    elif template == "vanguard":
        _band(kit, view)
        rows = [(label, value) for label, value in (("Date", view["date"]), ("To", view["recipient"]),
                                                     ("Re", view["subject"])) if value]
        if rows:
            table = kit.table([GUTTER, kit.width - GUTTER], rows=len(rows))
            for n, (row, (label, value)) in enumerate(zip(table.rows, rows, strict=True)):
                left, right = row.cells
                kit.cell_style(left, pad=(12 if not n else 3, 6, 3, 0))
                kit.cell_style(right, pad=(12 if not n else 3, 0, 3, 8))
                kit.first_paragraph(left, "Section Heading").add_run(label)
                left.paragraphs[0].paragraph_format.space_before = Pt(0)
                left.paragraphs[0].paragraph_format.keep_with_next = False
                cell = kit.first_paragraph(right, "Recipient")
                for i, line in enumerate(value if isinstance(value, list) else [value]):
                    if i:
                        cell.add_run().add_break(WD_BREAK.LINE)
                    cell.add_run(line).bold = label == "To" and i == 0
        start = len(kit.doc.paragraphs)
        _letter_text(kit, kit.doc, view, meta=False)
        for paragraph in kit.doc.paragraphs[start:]:
            paragraph.paragraph_format.left_indent = GUTTER + Pt(8)
    else:
        _header(kit, kit.doc, view)
        _letter_text(kit, kit.doc, view)
    kit.footer(view["name"])
    return kit.save(metadata)

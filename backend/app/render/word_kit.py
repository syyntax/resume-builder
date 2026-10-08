"""The DOCX style kit (DESIGN.md §11.6): named Word styles in the theme's colors, real bullet numbering, links,
borders, shading and the page-2+ footer, for the per-template builders in word.py.

python-docx covers styles and most formatting; borders, shading, hyperlinks, fields and character spacing are
written as XML here, each inserted where the schema puts it (Word refuses files with elements out of order).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.render.theme import DENSITIES, TEMPLATES, Theme

# Typefaces most recipients already have, per pairing (heading, body): the bundled web fonts aren't embedded.
SANS, SERIF = "Calibri", "Georgia"
WORD_FONTS: dict[str, tuple[str, str]] = {
    "inter": (SANS, SANS), "plex": (SANS, SANS), "source": (SANS, SERIF), "source-sans": (SANS, SANS),
    "garamond": (SERIF, SERIF), "lato": (SANS, SERIF),
}
PAPER = {"letter": (Inches(8.5), Inches(11)), "a4": (Emu(7560310), Emu(10692130))}
INK, TEXT, MUTED, RULE = "0F172A", "1F2937", "64748B", "D6DDE6"

# Child order of w:pPr, w:rPr, w:tcPr and w:tblPr in the schema (ECMA-376), for inserting elements in place.
_PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
        "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
        "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
        "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
        "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
_RPR = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
        "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
        "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em",
        "lang", "eastAsianLayout", "specVanish", "oMath"]
_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
         "tcFitText", "vAlign", "hideMark"]
_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
          "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def _place(parent: Any, child: Any, order: list[str]) -> Any:
    """Put ``child`` into ``parent`` where the schema expects it, replacing one with the same tag."""
    name = child.tag.split("}")[1]
    for old in parent.findall(qn(f"w:{name}")):
        parent.remove(old)
    successors = {qn(f"w:{n}") for n in order[order.index(name) + 1:]}
    following = next((c for c in parent if c.tag in successors), None)
    if following is None:
        parent.append(child)
    else:
        following.addprevious(child)
    return child


def _el(tag: str, **attrs: str) -> Any:
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(qn(f"w:{key}"), value)
    return element


def mix(hex_color: str, amount: float) -> str:
    """The color mixed with white (amount 1 is the color), as RRGGBB."""
    channels = (int(hex_color[i:i + 2], 16) for i in (1, 3, 5))
    return "".join(f"{round(255 - (255 - c) * amount):02X}" for c in channels)


def twips(length: Emu | int) -> str:
    return str(round(Emu(length).pt * 20))


@dataclass(frozen=True)
class Look:
    """How a template dresses the shared styles."""

    name_size: float = 26
    name_color: str = "accent"  # accent | ink | white
    name_caps: bool = False
    name_center: bool = False
    headline_color: str = "475569"
    headline_italic: bool = False
    headline_caps: bool = False
    heading_color: str = "accent"
    heading_size: float = 8.5
    heading_font: str = "heading"
    heading_rule: str | None = RULE  # bottom border color, or None
    org_color: str = "accent"
    bullet: str = "•"
    contact_center: bool = False


LOOKS: dict[str, Look] = {
    "meridian": Look(),
    "atlas": Look(name_size=25, name_color="ink", headline_color="accent", heading_color="ink", heading_rule=None,
                  bullet="–"),
    "ledger": Look(name_size=22, name_color="ink", name_caps=True, name_center=True, headline_italic=True,
                   heading_size=11, heading_rule="accent-light", org_color="ink", contact_center=True),
    "vanguard": Look(name_size=28, name_color="white", headline_color="white", headline_caps=True,
                     heading_rule=None, bullet="›"),
}


class Kit:
    """A document with the template's page setup, styles and bullet numbering."""

    def __init__(self, template: str, theme: Theme) -> None:
        self.template, self.theme, self.look = template, theme, LOOKS[template]
        self.info = TEMPLATES[template]
        self.accent = theme.accent[1:]
        self.heading_font, self.body_font = WORD_FONTS[theme.fonts]
        density = DENSITIES[theme.density]
        self.size = max(10.0, round(density.size_pt * 2) / 2)  # Word fonts run smaller than the web ones
        self.space = density.space
        self.doc = Document()
        top, sides, bottom = self.info.margins[theme.margins]
        section = self.doc.sections[0]
        section.page_width, section.page_height = PAPER[theme.paper]
        section.top_margin, section.bottom_margin = Inches(max(top, 0.5)), Inches(max(bottom, 0.5))
        section.left_margin = section.right_margin = Inches(sides)
        section.start_type = WD_SECTION_START.NEW_PAGE
        self.width = section.page_width - section.left_margin - section.right_margin
        self._styles()
        self.bullet_num = self._numbering()

    def color(self, name: str) -> str:
        return {"accent": self.accent, "ink": INK, "white": "FFFFFF", "accent-light": mix(self.theme.accent, 0.3),
                "tint": mix(self.theme.accent, 0.08)}.get(name, name)

    # ---- styles ------------------------------------------------------------------------------------------

    def _style(self, name: str, *, base: str = "Normal", size: float | None = None, bold: bool | None = None,
               italic: bool | None = None, color: str | None = TEXT, font: str | None = None, caps: bool = False,
               before: float = 0, after: float = 0, align: Any = None, keep_next: bool = False,
               tracking: float = 0, outline: int | None = None, kind: Any = WD_STYLE_TYPE.PARAGRAPH) -> Any:
        style = self.doc.styles.add_style(name, kind)
        style.base_style = self.doc.styles[base] if kind == WD_STYLE_TYPE.PARAGRAPH else None
        style.quick_style = True
        f = style.font  # a character style sets only what it's given, so it inherits the rest from its paragraph
        paragraph = kind == WD_STYLE_TYPE.PARAGRAPH
        if size or paragraph:
            f.size = Pt(size or self.size)
        if bold is not None or paragraph:
            f.bold = bool(bold)
        if italic is not None or paragraph:
            f.italic = bool(italic)
        if caps:
            f.all_caps = True
        if color:
            f.color.rgb = RGBColor.from_string(self.color(color))
        if font:
            self._fonts(style.element.get_or_add_rPr(), font)
        if tracking:
            _place(style.element.get_or_add_rPr(), _el("w:spacing", val=str(round(tracking * 20))), _RPR)
        if kind == WD_STYLE_TYPE.PARAGRAPH:
            pf = style.paragraph_format
            pf.space_before, pf.space_after = Pt(before * self.space), Pt(after * self.space)
            pf.keep_with_next = keep_next
            if align is not None:
                pf.alignment = align
            if outline is not None:
                _place(style.element.get_or_add_pPr(), _el("w:outlineLvl", val=str(outline)), _PPR)
        return style

    @staticmethod
    def _fonts(rpr: Any, font: str) -> None:
        fonts = rpr.get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
            fonts.set(qn(f"w:{attr}"), font)

    def _styles(self) -> None:
        look, styles = self.look, self.doc.styles
        normal = styles["Normal"]
        normal.font.size, normal.font.color.rgb = Pt(self.size), RGBColor.from_string(TEXT)
        self._fonts(normal.element.get_or_add_rPr(), self.body_font)
        pf = normal.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.MULTIPLE, 1.1
        heading = self.heading_font if look.heading_font == "heading" else self.body_font
        center = WD_ALIGN_PARAGRAPH.CENTER
        self._style("Name", size=look.name_size, bold=not look.name_caps, color=look.name_color, font=self.heading_font,
                    caps=look.name_caps, tracking=1.5 if look.name_caps else 0, outline=0,
                    align=center if look.name_center else None)
        self._style("Headline", size=self.size * 1.15, color=look.headline_color, italic=look.headline_italic,
                    caps=look.headline_caps, bold=look.headline_caps or self.template == "atlas", before=3,
                    tracking=1 if look.headline_caps else 0, align=center if look.name_center else None)
        self._style("Contact", size=self.size * 0.92, color="334155" if look.name_color != "white" else "white",
                    before=6, align=center if look.contact_center else None)
        style = self._style("Section Heading", size=look.heading_size, bold=True, color=look.heading_color,
                            font=heading, caps=True, tracking=0.6, before=11, after=4, keep_next=True, outline=1)
        if look.heading_rule:
            self.border(style.element.get_or_add_pPr(), "bottom", self.color(look.heading_rule), size=4, space=2)
        self._style("Role Title", size=self.size * 1.05, bold=True, color=INK, font=self.heading_font, keep_next=True)
        self._style("Org", bold=True, color=look.org_color, keep_next=True, after=1,
                    size=self.size * (1.2 if self.template == "ledger" else 1), font=self.heading_font
                    if self.template == "ledger" else None)
        self._style("Dates", kind=WD_STYLE_TYPE.CHARACTER, size=self.size * 0.92, color=MUTED,
                    italic=self.template == "ledger", bold=False)
        self._style("Muted", kind=WD_STYLE_TYPE.CHARACTER, color=MUTED)
        bullet = self._style("Bullet", after=2)
        bullet.paragraph_format.left_indent = Inches(0.2)
        bullet.paragraph_format.first_line_indent = Inches(-0.2)
        self._style("Body", after=2)
        self._style("Letter Body", size=max(11.0, self.size + 1), after=8)
        self._style("Recipient", size=max(11.0, self.size + 1))
        self._style("Signature", size=max(12.0, self.size + 2), bold=self.template != "ledger", font=self.heading_font,
                    color="accent" if self.template in ("meridian", "vanguard") else INK, before=14)
        self._style("Subject", size=self.size, bold=True, color="accent", font=heading, caps=True, tracking=0.6,
                    before=14)
        self._style("Sidebar", size=self.size * 0.9)
        link = self._style("Link", kind=WD_STYLE_TYPE.CHARACTER, color=None)
        link.font.underline = False

    # ---- numbering ---------------------------------------------------------------------------------------

    def _numbering(self) -> int:
        """A bullet list definition in the accent color, attached to the Bullet style; returns its numId."""
        numbering = self.doc.part.numbering_part.element
        ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        abstract_id = max(ids, default=0) + 1
        abstract = _el("w:abstractNum", abstractNumId=str(abstract_id))
        abstract.append(_el("w:multiLevelType", val="singleLevel"))
        level = _el("w:lvl", ilvl="0")
        for child in (_el("w:start", val="1"), _el("w:numFmt", val="bullet"), _el("w:lvlText", val=self.look.bullet),
                      _el("w:lvlJc", val="left")):
            level.append(child)
        ppr = _el("w:pPr")
        ppr.append(_el("w:ind", left="288", hanging="288"))
        level.append(ppr)
        rpr = _el("w:rPr")
        rpr.append(_el("w:rFonts", ascii=self.body_font, hAnsi=self.body_font, hint="default"))
        rpr.append(_el("w:b"))
        rpr.append(_el("w:color", val=self.accent))
        level.append(rpr)
        abstract.append(level)
        # abstractNum elements come before num elements
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        num = numbering.add_num(abstract_id)
        style_ppr = self.doc.styles["Bullet"].element.get_or_add_pPr()
        num_pr = style_ppr.get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num.numId
        return num.numId

    # ---- building blocks ---------------------------------------------------------------------------------

    @staticmethod
    def border(ppr: Any, side: str, color: str, *, size: int = 4, space: int = 1, kind: str = "single") -> None:
        borders = ppr.find(qn("w:pBdr"))
        if borders is None:
            borders = _place(ppr, _el("w:pBdr"), _PPR)
        edge = _el(f"w:{side}", val=kind, sz=str(size), space=str(space), color=color)
        _place(borders, edge, ["top", "left", "bottom", "right", "between", "bar"])

    def para(self, container: Any, style: str, text: str = "", *, inline: bool = True) -> Paragraph:
        paragraph = container.add_paragraph(style=style)
        if text:
            self.runs(paragraph, text) if inline else paragraph.add_run(text)
        return paragraph

    @staticmethod
    def runs(paragraph: Paragraph, text: str, **fmt: Any) -> None:
        """Text with **bold** and *italic* as runs (the only markup the documents carry)."""
        for part in re.split(r"(\*\*\S(?:.*?\S)?\*\*|(?<![*\w])\*\S(?:.*?\S)?\*(?![*\w]))", text):
            if not part:
                continue
            run = paragraph.add_run()
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                run.text, run.bold = part[2:-2], True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                run.text, run.italic = part[1:-1], True
            else:
                run.text = part
            for key, value in fmt.items():
                setattr(run, key, value)

    def right_tab(self, paragraph: Paragraph, width: Emu | int) -> None:
        paragraph.paragraph_format.tab_stops.add_tab_stop(Emu(width), WD_TAB_ALIGNMENT.RIGHT)

    def row(self, container: Any, style: str, left: str, right: str, width: Emu | int, *,
            left_style: str | None = None) -> Paragraph:
        """A line with text on the left and, right-aligned on a tab stop, dates or a place."""
        paragraph = container.add_paragraph(style=style)
        run = paragraph.add_run(left)
        if left_style:
            run.style = self.doc.styles[left_style]
        if right:
            self.right_tab(paragraph, width)
            paragraph.add_run("\t" + right, style=self.doc.styles["Dates"])
        return paragraph

    def hyperlink(self, paragraph: Paragraph, text: str, url: str) -> None:
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = _el("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = paragraph.add_run(text, style=self.doc.styles["Link"])
        run._r.getparent().remove(run._r)  # noqa: SLF001  (python-docx has no hyperlink API)
        link.append(run._r)  # noqa: SLF001
        paragraph._p.append(link)  # noqa: SLF001

    def contact(self, container: Any, items: list[dict[str, Any]], *, style: str = "Contact",
                stacked: bool = False, color: str | None = None) -> None:
        """Contact details, links included: on one line separated by dots, or one per line."""
        paragraph = None
        for i, item in enumerate(items):
            if stacked or paragraph is None:
                paragraph = container.add_paragraph(style=style)
                if stacked and i:
                    paragraph.paragraph_format.space_before = Pt(2)
            elif i:
                separator = paragraph.add_run("  ·  ")
                separator.font.color.rgb = RGBColor.from_string(color or self.accent)
            if item.get("href") and item["href"].startswith(("https://", "mailto:", "tel:")):
                self.hyperlink(paragraph, item["text"], item["href"])
            else:
                paragraph.add_run(item["text"])

    def bullets(self, container: Any, items: list[dict[str, Any]]) -> None:
        for bullet in items:
            self.para(container, "Bullet", bullet.get("text", ""))

    def footer(self, name: str) -> None:
        """'Name · Page 2 of 3' from page 2 on; nothing on page 1 (§11.2)."""
        section = self.doc.sections[0]
        section.different_first_page_header_footer = True
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for part in (f"{name} · Page ", "PAGE", " of ", "NUMPAGES"):
            run = paragraph.add_run("1" if part in ("PAGE", "NUMPAGES") else part)
            run.font.size, run.font.color.rgb = Pt(7.5), RGBColor.from_string(MUTED)
            if part in ("PAGE", "NUMPAGES"):  # a field whose result Word updates: wrap the run in it
                field = _el("w:fldSimple", instr=f" {part} ")
                run._r.addprevious(field)  # noqa: SLF001
                field.append(run._r)  # noqa: SLF001

    # ---- tables (Atlas's sidebar, Vanguard's band and gutter) -------------------------------------------------

    def table(self, widths: list[Emu | int], rows: int = 1) -> Table:
        """A borderless table with fixed column widths."""
        table = self.doc.add_table(rows=rows, cols=len(widths))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        tbl_pr = table._tbl.tblPr  # noqa: SLF001
        _place(tbl_pr, _el("w:tblW", w=twips(sum(widths, Emu(0))), type="dxa"), _TBLPR)
        borders = _el("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(_el(f"w:{side}", val="nil"))
        _place(tbl_pr, borders, _TBLPR)
        _place(tbl_pr, _el("w:tblLayout", type="fixed"), _TBLPR)
        margins = _el("w:tblCellMar")
        for side, value in (("top", "0"), ("left", "0"), ("bottom", "0"), ("right", "0")):
            margins.append(_el(f"w:{side}", w=value, type="dxa"))
        _place(tbl_pr, margins, _TBLPR)
        for col, width in zip(table.columns, widths, strict=True):
            col.width = width
        for row in table.rows:
            for cell, width in zip(row.cells, widths, strict=True):
                cell.width = width
        return table

    def cell_style(self, cell: _Cell, *, fill: str | None = None, pad: tuple[float, float, float, float] | None = None,
                   top_border: str | None = None) -> None:
        """Shading, padding (top, right, bottom, left, in points) and a top border for one cell."""
        tc_pr = cell._tc.get_or_add_tcPr()  # noqa: SLF001
        if top_border:
            borders = _el("w:tcBorders")
            borders.append(_el("w:top", val="single", sz="4", space="0", color=top_border))
            _place(tc_pr, borders, _TCPR)
        if fill:
            _place(tc_pr, _el("w:shd", val="clear", color="auto", fill=fill), _TCPR)
        if pad:
            margins = _el("w:tcMar")
            for side, value in zip(("top", "right", "bottom", "left"), pad, strict=True):
                margins.append(_el(f"w:{side}", w=str(round(value * 20)), type="dxa"))
            # schema order inside tcMar: top, start/left, bottom, end/right
            ordered = sorted(margins, key=lambda e: ["top", "left", "bottom", "right"].index(e.tag.split("}")[1]))
            for element in list(margins):
                margins.remove(element)
            for element in ordered:
                margins.append(element)
            _place(tc_pr, margins, _TCPR)

    @staticmethod
    def first_paragraph(cell: _Cell, style: str) -> Paragraph:
        """A new cell holds one empty paragraph: use it for the first block rather than leaving a blank line."""
        paragraph = cell.paragraphs[0]
        paragraph.style = style
        return paragraph

    # ---- output --------------------------------------------------------------------------------------------

    def save(self, metadata: dict[str, str]) -> bytes:
        body = self.doc.element.body
        last = body[-2] if body[-1].tag == qn("w:sectPr") else body[-1]
        if last.tag == qn("w:tbl"):  # Word wants a paragraph after a final table: make it as small as possible
            end = self.doc.add_paragraph()
            end.paragraph_format.space_before = end.paragraph_format.space_after = Pt(0)
            end.paragraph_format.line_spacing = Pt(1)
            end.add_run().font.size = Pt(1)
        props = self.doc.core_properties
        props.title, props.author = metadata.get("Title", ""), metadata.get("Author", "")
        props.subject, props.keywords = metadata.get("Subject", ""), metadata.get("Keywords", "")
        props.comments = "Made with Resume Tailor"
        out = BytesIO()
        self.doc.save(out)
        return out.getvalue()

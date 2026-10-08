"""DOCX export (DESIGN.md §11.6, EXP-2): every template's résumé and cover letter as an editable Word document."""

from __future__ import annotations

import copy
import io
import re
import zipfile
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi import FastAPI
from lxml import etree

from app.pipeline.export import markdown_to_text, to_markdown
from app.render.samples import LETTER, RESUME, long_resume
from app.render.theme import TEMPLATES, Theme
from app.render.word import build_letter_docx, build_resume_docx
from app.render.word_kit import _PPR, _RPR, _TBLPR, _TCPR, WORD_FONTS
from tests.conftest import POSTING, Api, drain

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
META = {"Title": "Jordan Rivera – Staff Platform Engineer – Globex", "Author": "Jordan Rivera",
        "Subject": "Résumé for Staff Platform Engineer at Globex", "Keywords": "Kubernetes, Terraform"}


def theme_for(template: str, **options: Any) -> Theme:
    info = TEMPLATES[template]
    return Theme(**{"accent": info.accent, "fonts": info.fonts, "contact_icons": info.contact_icons, **options})


def body_text(data: bytes) -> str:
    """The document's text in reading order, tables included, whitespace collapsed."""
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    words = []
    for paragraph in root.iter(f"{W}p"):
        words.append("".join(t.text or "" for t in paragraph.iter(f"{W}t", f"{W}tab")
                             if t.tag == f"{W}t") + " ")
    return " ".join("".join(words).split())


def assert_in_order(text: str, order: list[str]) -> None:
    positions = [text.find(t) for t in order]
    assert -1 not in positions and positions == sorted(positions), dict(zip(order, positions, strict=True))


def assert_schema_order(data: bytes) -> None:
    """Word refuses a file whose property elements are out of schema order: check every part we write."""
    archive = zipfile.ZipFile(io.BytesIO(data))
    for part in archive.namelist():
        if not part.startswith("word/") or not part.endswith(".xml"):
            continue
        root = etree.fromstring(archive.read(part))
        for tag, order in (("pPr", _PPR), ("rPr", _RPR), ("tcPr", _TCPR), ("tblPr", _TBLPR)):
            for element in root.iter(f"{W}{tag}"):
                names = [c.tag.split("}")[1] for c in element if c.tag.startswith(W)]
                known = [n for n in names if n in order]
                assert known == sorted(known, key=order.index), (part, tag, names)


@pytest.mark.parametrize("template", list(TEMPLATES))
def test_resume_docx(template: str) -> None:
    resume = copy.deepcopy(RESUME)
    resume["summary"]["text"] = "Shipped **fast** and *carefully*; <b>not html</b>."
    data = build_resume_docx(resume, template, theme_for(template), META)
    assert_schema_order(data)
    doc = Document(io.BytesIO(data))
    text = body_text(data)

    main = ["Jordan Rivera", "Senior Platform Engineer", "Shipped fast and carefully; <b>not html</b>.",
            "Senior Software Engineer, Platform", "Acme Payments", "Led migration of 40 payment services", "Initech",
            "Software Engineer II", "Earlier", "tf-guard"]
    side = ["jordan@example.com", "Languages", "Go, Python, TypeScript, Bash", "Certifications",
            "Certified Kubernetes Administrator (CKA)", "Education", "B.S. Computer Science"]
    if template == "ledger":  # the employer, then the title
        main[3], main[4] = main[4], main[3]
    # Atlas's sidebar is the table's first cell; everything else reads top to bottom
    assert_in_order(text, side + main if template == "atlas" else main[:2] + side[:1] + main[2:] + side[1:])
    assert "**" not in text

    used = {p.style.name for p in doc.paragraphs} | {p.style.name for t in doc.tables for row in t.rows
                                                      for cell in row.cells for p in cell.paragraphs}
    assert {"Name", "Headline", "Section Heading", "Bullet", "Body"} <= used
    bullets = [p for t in [doc] + [c for t in doc.tables for r in t.rows for c in r.cells] for p in t.paragraphs
               if p.style.name == "Bullet"]
    assert len(bullets) == 9 and bullets[0].text.startswith("Led migration")
    # real Word numbering on the Bullet style, in the accent color
    num_id = doc.styles["Bullet"].element.pPr.numPr.numId.val
    numbering = doc.part.numbering_part.element
    num = next(n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == num_id)
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract_id)
    assert abstract.find(f"{qn('w:lvl')}/{qn('w:numFmt')}").get(qn("w:val")) == "bullet"
    color = abstract.find(f"{qn('w:lvl')}/{qn('w:rPr')}/{qn('w:color')}").get(qn("w:val"))
    assert color == TEMPLATES[template].accent[1:]
    bold = [r for p in doc.paragraphs + [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
            for r in p.runs if r.text == "fast"]
    assert bold and bold[0].bold

    links = {rel.target_ref for rel in doc.part.rels.values() if rel.reltype.endswith("/hyperlink")}
    assert {"mailto:jordan@example.com", "https://linkedin.com/in/jrivera"} <= links
    section = doc.sections[0]
    assert (round(section.page_width.inches, 2), round(section.page_height.inches, 2)) == (8.5, 11)
    assert section.different_first_page_header_footer  # no footer on page 1
    footer_xml = etree.tostring(section.footer._element, encoding="unicode")  # noqa: SLF001
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml and "Jordan Rivera · Page" in footer_xml
    fonts = doc.styles["Normal"].element.rPr.rFonts
    assert fonts.get(qn("w:ascii")) == WORD_FONTS[TEMPLATES[template].fonts][1]
    props = doc.core_properties
    assert (props.title, props.author, props.keywords) == (META["Title"], META["Author"], META["Keywords"])


@pytest.mark.parametrize("template", list(TEMPLATES))
def test_cover_letter_docx(template: str) -> None:
    letter = copy.deepcopy(LETTER)
    letter["highlights"] = [{"value": "22 teams", "label": "on the platform I designed", "sourceIds": []}]
    data = build_letter_docx(letter, template, theme_for(template), {"Title": "Letter"})
    assert_schema_order(data)
    text = body_text(data)
    subject = "Staff Platform Engineer" if template == "vanguard" else "Re: Staff Platform Engineer"
    assert_in_order(text, ["Jordan Rivera", "October 6, 2026", "Priya Shah", "Director of Platform Engineering",
                           subject, "Dear Priya Shah,", "Globex is looking", "Sincerely,"])
    assert ("22 teams" in text) == (template == "atlas")
    assert Document(io.BytesIO(data)).paragraphs[-1].text in ("", "Jordan Rivera")


def test_docx_follows_the_theme() -> None:
    data = build_resume_docx(long_resume(), "meridian", Theme(paper="a4", fonts="garamond", accent="#7A1F3D",
                                                              density="comfortable", margins="narrow"), META)
    assert_schema_order(data)
    doc = Document(io.BytesIO(data))
    section = doc.sections[0]
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (210, 297)
    assert round(section.left_margin.inches, 2) == 0.5
    assert str(doc.styles["Name"].font.color.rgb) == "7A1F3D"
    assert doc.styles["Normal"].element.rPr.rFonts.get(qn("w:ascii")) == "Georgia"
    assert doc.styles["Normal"].font.size.pt == 10


def test_plain_text_export() -> None:
    resume = copy.deepcopy(RESUME)
    resume["summary"]["text"] = "Shipped **fast** with C# and 3 * 4 = 12 checks."
    text = markdown_to_text(to_markdown(resume))
    assert text.startswith("Jordan Rivera\nSenior Platform Engineer\n")
    assert "\nSUMMARY\n" in text and "\nEXPERIENCE\n" in text
    assert "Shipped fast with C# and 3 * 4 = 12 checks." in text
    assert "- Led migration of 40 payment services" in text
    assert not re.search(r"(^|\s)#|\\", text)


def test_export_formats(app: FastAPI, keyed_user: Api, example: dict[str, Any]) -> None:
    master_id = keyed_user.post("/api/master-resumes", json=example).json()["master"]["id"]
    app_id = keyed_user.post("/api/applications", json={"masterId": master_id, "job": {"text": POSTING}}).json()["id"]
    drain(app)
    keyed_user.put(f"/api/applications/{app_id}/design", json={"template": "vanguard", "theme": {"fonts": "lato"}})
    for doc, kind in (("resume", "Resume"), ("cover-letter", "Cover-Letter")):
        docx = keyed_user.get(f"/api/applications/{app_id}/export", params={"doc": doc, "format": "docx"})
        assert docx.status_code == 200, docx.text
        assert docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert docx.headers["content-disposition"].endswith(f'_{kind}.docx"')
        parsed = Document(io.BytesIO(docx.content))
        assert parsed.styles["Normal"].element.rPr.rFonts.get(qn("w:ascii")) == "Georgia"  # Lato / Merriweather
        assert len(parsed.tables) >= 1  # Vanguard's band
        txt = keyed_user.get(f"/api/applications/{app_id}/export", params={"doc": doc, "format": "txt"})
        assert txt.status_code == 200 and txt.headers["content-type"].startswith("text/plain")
        assert "**" not in txt.text and txt.text.strip()
    packet = keyed_user.get(f"/api/applications/{app_id}/export", params={"doc": "packet", "format": "docx"})
    assert packet.status_code == 422

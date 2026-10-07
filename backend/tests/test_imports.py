"""Résumé import (DESIGN.md §10): file handling, extraction, conversion, fidelity, and the API flow."""

from __future__ import annotations

import asyncio
import io
import zipfile
from typing import Any

import pytest
from fastapi import FastAPI

from app.ai.fake import FakeLLM
from app.ai.import_schema import MasterImportDraft
from app.ai.llm import AnthropicLLM, LLMError, LLMRequest, LLMResult
from app.importers import extract as extract_module
from app.importers import fidelity
from app.importers.convert import convert, finalize, locate, parse_date
from app.importers.extract import extract
from app.importers.sniff import ImportRejected, sniff
from tests import corpus
from tests.conftest import USER_PASSWORD, Api, drain

PDF = {"content-type": "application/pdf"}


def run(coro: Any) -> Any:
    return asyncio.run(coro)


def upload(user: Api, data: bytes, name: str = "resume.pdf") -> Any:
    return user.post(f"/api/imports?filename={name}", content=data,
                     headers={"content-type": "application/octet-stream"})


def imported(app: FastAPI, user: Api, data: bytes, name: str = "resume.pdf") -> dict[str, Any]:
    response = upload(user, data, name)
    assert response.status_code == 202, response.text
    drain(app)
    detail = user.get(f"/api/imports/{response.json()['id']}").json()
    assert detail["status"] == "ready", detail
    return detail


# ---- sniffing and extraction ---------------------------------------------------------------------------


def test_files_are_identified_by_content_not_name() -> None:
    assert sniff(corpus.text_pdf()) == "pdf"
    assert sniff(corpus.full_docx()) == "docx"
    old_word = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 99
    for data, code in [(b"", "empty"), (b"hello world", "unsupported"), (old_word, "unsupported"),
                       (corpus.docm(), "macros"), (corpus.zip_bomb(), "too_large"),
                       (b"%PDF-" + b"0" * 6_000_000, "too_large")]:
        with pytest.raises(ImportRejected) as raised:
            sniff(data)
        assert raised.value.code == code
    xlsx = io.BytesIO()
    with zipfile.ZipFile(xlsx, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
    with pytest.raises(ImportRejected, match="spreadsheets"):
        sniff(xlsx.getvalue())


def test_pdf_extraction() -> None:
    text = run(extract("pdf", corpus.text_pdf()))
    assert text.pages == 1 and text.has_text_layer
    assert "Built a canary release pipeline" in text.text

    scan = run(extract("pdf", corpus.scanned_pdf()))
    assert scan.pages == 1 and not scan.has_text_layer and scan.text == ""

    linkedin = run(extract("pdf", corpus.linkedin_pdf()))
    assert linkedin.pages == 2 and "Top Skills" in linkedin.text


@pytest.mark.parametrize(("make", "code"), [
    (corpus.encrypted_pdf, "encrypted"), (corpus.long_pdf, "too_many_pages"), (corpus.lying_zip_bomb, "unreadable"),
    (corpus.docx_with_dtd, "unsafe"), (lambda: b"%PDF-1.7\n garbage", "unreadable"),
])
def test_hostile_or_unusable_files_are_rejected(make: Any, code: str) -> None:
    data = make()
    with pytest.raises(ImportRejected) as raised:
        run(extract("docx" if data.startswith(b"PK") else "pdf", data))
    assert raised.value.code == code


def test_docx_extraction_reads_every_part_and_sanitizes_html() -> None:
    result = run(extract("docx", corpus.full_docx()))
    text = result.text
    assert text.startswith("[Header]\nJordan Rivera\n")  # contact details often live in the header
    assert "# Experience" in text and "- Led migration of 40 payment services" in text  # headings, list items
    assert "Cloud & Infrastructure | AWS, Kubernetes, Terraform" in text  # table
    assert "# Certifications\nCertified Kubernetes Administrator - CNCF (2023)" in text  # text box, in place
    assert "<https://github.com/jrivera/tf-guard>" in text  # hyperlink target
    assert text.endswith("[Footer]\nRésumé of Jordan Rivera – references on request")
    assert "<h1>Experience</h1>" in (result.html or "") and "<ul><li>" in (result.html or "")


def test_docx_preview_html_is_sanitized() -> None:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = Document()
    paragraph = document.add_paragraph("Click ")
    r_id = paragraph.part.relate_to("javascript:alert(1)", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    paragraph._p.append(parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}"><w:r><w:t>here</w:t></w:r>'
                                  "</w:hyperlink>"))
    document.add_paragraph("<script>alert(2)</script>")
    buffer = io.BytesIO()
    document.save(buffer)
    html = run(extract("docx", buffer.getvalue())).html or ""
    assert "javascript:" not in html and "<script" not in html and "&lt;script&gt;" in html


def test_the_worker_gets_no_secrets_and_failures_are_contained(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("APP_SECRET_KEY", "x" * 40)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-secret")
    env = extract_module._worker_env()
    assert not any("KEY" in name or "SECRET" in name for name in env)
    # A worker that dies (here: a module that doesn't exist) is reported as an unreadable file, not a crash.
    original = asyncio.create_subprocess_exec

    async def broken(*args: Any, **kwargs: Any) -> Any:
        return await original(args[0], "-m", "app.importers.no_such_worker", **kwargs)

    monkeypatch.setattr(extract_module.asyncio, "create_subprocess_exec", broken)
    with pytest.raises(ImportRejected) as raised:
        run(extract("pdf", corpus.text_pdf()))
    assert raised.value.code == "unreadable"


# ---- conversion and fidelity -----------------------------------------------------------------------------


@pytest.mark.parametrize(("text", "expected"), [
    ("2021", "2021"), ("2021-3", "2021-03"), ("2021-03-09", "2021-03-09"), ("03/2021", "2021-03"),
    ("Mar 2021", "2021-03"), ("Sept. 2019", "2019-09"), ("September, 2019", "2019-09"), ("Spring 2019", None),
    ("2021-13", None), ("", None),
])
def test_dates_are_read_leniently(text: str, expected: str | None) -> None:
    assert parse_date(text) == expected


def _draft(**overrides: Any) -> MasterImportDraft:
    empty_location = {"city": "", "region": "", "countryCode": ""}
    base: dict[str, Any] = {
        "basics": {"name": "Sam Lee", "label": "", "email": "", "phone": "", "url": "", "location": empty_location,
                   "profiles": []},
        "headline": "", "summary": "", "work": [], "projects": [], "education": [], "certifications": [],
        "skills": [], "inferredSkills": [], "awards": [], "publications": [], "volunteer": [], "languages": [],
        "interests": [], "customSections": [], "importNotes": [],
    }
    return MasterImportDraft.model_validate({**base, **overrides})


def _role(**overrides: Any) -> dict[str, Any]:
    return {"company": "Acme", "position": "Engineer", "location": "", "startDate": "2020-01", "endDate": "",
            "current": True, "summary": "", "highlights": ["Did things."], "tags": [], **overrides}


def test_conversion_fills_gaps_visibly_and_keeps_notes_pointing_at_the_right_item() -> None:
    draft = _draft(
        headline="Engineer",
        work=[_role(company="", position="", highlights=[], summary=""),  # empty: dropped
              _role(company="", highlights=["Shipped the thing."]),  # no company: placeholder + note
              _role(company="Initech", startDate="Spring 2019", current=False, endDate="")],
        skills=[{"name": "Go", "category": "Languages", "level": "", "aliases": ["golang"]},
                {"name": "go", "category": "", "level": "", "aliases": ["Go lang"]}],
        inferredSkills=[{"name": "Kafka", "category": "Data", "evidence": "built Kafka pipelines"},
                        {"name": "Go", "category": "", "evidence": "dup of a listed skill"}],
        importNotes=[{"path": "work[2].endDate", "severity": "warning", "message": "End date unclear."},
                     {"path": "headline", "severity": "info", "message": "Taken from the first line."}],
    )
    converted = convert(draft, label="Imported")
    document, notes = finalize(converted)
    work = document["work"]
    assert [w["company"] for w in work] == ["Unknown company", "Initech"]
    assert "startDate" not in work[1] and "endDate" not in work[1]
    assert document["headlines"][0]["text"] == "Engineer"
    assert document["skills"] == [{"id": "go", "name": "Go", "category": "Languages", "aliases": ["golang", "Go lang"]}]
    assert [s["name"] for s in converted.inferred_skills] == ["Kafka"]

    by_path = {n.path: n for n in notes}
    assert by_path["work[1].endDate"].message == "End date unclear."  # model note, re-indexed past the dropped role
    assert by_path["headlines[0].text"].severity == "info"
    assert "Unknown company" in by_path["work[0].company"].message
    assert "Spring 2019" in by_path["work[1].startDate"].message
    assert locate(document, "work[1].endDate") == (work[1]["id"], "endDate")
    assert locate(document, "work[0].highlights[0]") == (work[0]["highlights"][0]["id"], None)
    assert locate(document, "basics.phone") == ("basics", "phone")
    assert locate(document, "skills") == (None, None)


def test_values_that_fail_validation_are_removed_with_a_note() -> None:
    draft = _draft(work=[_role(startDate="2022-01", endDate="2020-01", current=False)])
    document, notes = finalize(convert(draft, label="x"))
    assert "endDate" not in document["work"][0]
    assert any("Removed endDate" in n.message for n in notes)


def test_fidelity_tolerates_extraction_noise_but_flags_changed_facts() -> None:
    source = ("Résumé Writer • Led migra-\ntion of 40 payment services from EC2 to EKS,\n  "
              "cutting p99 latency 45% and saving $380K/yr. Kuber netes.")
    bullets = [
        "Led migration of 40 payment services from EC2 to EKS, cutting p99 latency 45% and saving $380K/yr.",
        "Led migration of 400 payment services from EC2 to EKS",  # a number changed
        "Led migration of 40 payment services from EC2 to GKE",  # a word changed
        "Architected a globally distributed edge platform",  # not in the document at all
        "Kubernetes",  # split in two by extraction: fine
    ]
    doc = {"basics": {"name": "Résumé Writer"}, "work": [{"id": "w", "company": "EC2", "position": "UX",
            "highlights": [{"id": f"w.b{i}", "text": t} for i, t in enumerate(bullets)]}]}
    report = fidelity.check(doc, source)
    mismatches = {m["targetId"]: m for m in report["mismatches"]}
    assert set(mismatches) == {"w.b1", "w.b2", "w.b3"}
    assert mismatches["w.b1"]["missing"] == ["400"] and mismatches["w.b2"]["missing"] == ["GKE"]
    assert mismatches["w.b3"]["score"] < 90
    assert report["checked"] == 7 and report["passed"] == 4  # "UX" is too short to prove anything
    assert fidelity.check(doc, "")["available"] is False


# ---- the import flow over the API ------------------------------------------------------------------------------


def _ids(node: Any) -> list[str]:
    if isinstance(node, list):
        return [i for n in node for i in _ids(n)]
    if isinstance(node, dict):
        own = [node["id"]] if isinstance(node.get("id"), str) else []
        return own + [i for v in node.values() for i in _ids(v)]
    return []


CORPUS = {"text.pdf": corpus.text_pdf, "two-column.pdf": corpus.two_column_pdf, "linkedin.pdf": corpus.linkedin_pdf,
          "resume.docx": corpus.full_docx}


@pytest.mark.parametrize("name", list(CORPUS))
def test_corpus_imports_with_high_fidelity(app: FastAPI, keyed_user: Api, name: str) -> None:
    detail = imported(app, keyed_user, CORPUS[name](), name)
    assert detail["fidelity"]["available"] and detail["fidelity"]["score"] >= 0.95
    bullets = [h["text"] for w in detail["document"]["work"] for h in w["highlights"]]
    assert all(b in bullets for b in corpus.WORK[0][5])
    assert detail["sourceKind"] == name.rsplit(".", 1)[1] and detail["costUsd"] > 0
    # Every flag points somewhere the review screen can show: an item in the document, or a skill to add.
    ids = {"basics", *_ids(detail["document"])}
    for flag in detail["flags"]:
        if flag["kind"] == "inferredSkill":
            assert flag["skill"]["name"]
        else:
            assert flag["targetId"] in ids and flag["path"]


def test_review_and_accept_creates_a_master_with_the_original_attached(app: FastAPI, keyed_user: Api) -> None:
    pdf = corpus.text_pdf()
    detail = imported(app, keyed_user, pdf, "Jordan Résumé.pdf")
    import_id = detail["id"]
    assert detail["sourceName"] == "Jordan Résumé.pdf" and detail["pages"] == 1 and detail["hasTextLayer"]
    assert {f["kind"] for f in detail["flags"]} >= {"note", "inferredSkill"}

    source = keyed_user.get(f"/api/imports/{import_id}/source")
    assert source.content == pdf and source.headers["content-type"] == "application/pdf"
    assert "sandbox" in source.headers["content-security-policy"]
    assert "filename*=UTF-8''Jordan%20R%C3%A9sum%C3%A9.pdf" in source.headers["content-disposition"]
    assert keyed_user.get(f"/api/imports/{import_id}/preview").json()["kind"] == "pdf"

    edited = detail["document"]
    edited["basics"]["label"] = "Staff Platform Engineer"
    first_flag = detail["flags"][0]["key"]
    saved = keyed_user.put(f"/api/imports/{import_id}/review", json={"document": edited,
                                                                       "resolved": {first_flag: "dismissed"}})
    assert saved.status_code == 200, saved.text
    again = keyed_user.get(f"/api/imports/{import_id}").json()
    assert again["document"]["basics"]["label"] == "Staff Platform Engineer"
    assert again["resolved"] == {first_flag: "dismissed"}
    assert again["draft"]["basics"].get("label") != "Staff Platform Engineer"  # the draft is kept for reverting

    accepted = keyed_user.post(f"/api/imports/{import_id}/accept",
                               json={"label": "Imported master", "document": edited})
    assert accepted.status_code == 201, accepted.text
    master, version = accepted.json()["master"], accepted.json()["version"]
    assert master["label"] == "Imported master" and version["sourceFile"]["name"] == "Jordan Résumé.pdf"
    assert keyed_user.get(f"/api/imports/{import_id}").json()["status"] == "accepted"
    assert keyed_user.post(f"/api/imports/{import_id}/accept",
                           json={"label": "Again", "document": edited}).status_code == 409

    content = keyed_user.get(f"/api/master-resumes/{master['id']}").json()["content"]
    assert content["basics"]["label"] == "Staff Platform Engineer" and content["meta"]["label"] == "Imported master"
    versions = keyed_user.get(f"/api/master-resumes/{master['id']}/versions").json()
    assert versions[0]["sourceFile"]["kind"] == "pdf"
    file_url = f"/api/master-resumes/{master['id']}/versions/{version['id']}/source"
    assert keyed_user.get(file_url).content == pdf

    # Discarding the import keeps the file: the master's version still uses it.
    assert keyed_user.delete(f"/api/imports/{import_id}").status_code == 204
    assert keyed_user.get(file_url).status_code == 200
    # Deleting the file keeps the version.
    assert keyed_user.delete(file_url).status_code == 204
    assert keyed_user.get(file_url).status_code == 404
    assert keyed_user.get(f"/api/master-resumes/{master['id']}/versions").json()[0]["sourceFile"] is None
    assert list((app.state.settings.data_dir / "uploads").rglob("*")) == [
        app.state.settings.data_dir / "uploads" / keyed_user.user["id"]]


def test_accept_validates_the_reviewed_document(app: FastAPI, keyed_user: Api) -> None:
    detail = imported(app, keyed_user, corpus.full_docx(), "cv.docx")
    broken = {**detail["document"], "work": [{"position": "No company"}], "projects": []}
    response = keyed_user.post(f"/api/imports/{detail['id']}/accept", json={"label": "x", "document": broken})
    assert response.status_code == 422
    assert any(e["path"] == "work[0].company" for e in response.json()["errors"])
    preview = keyed_user.get(f"/api/imports/{detail['id']}/preview").json()
    assert preview["kind"] == "docx" and "<h1>Experience</h1>" in preview["html"]


def test_pasted_text_and_scanned_pdfs(app: FastAPI, keyed_user: Api) -> None:
    response = keyed_user.post("/api/imports/text", json={"text": corpus.plain_text()})
    assert response.status_code == 202
    drain(app)
    detail = keyed_user.get(f"/api/imports/{response.json()['id']}").json()
    assert detail["status"] == "ready" and detail["sourceName"] == "Pasted text" and detail["fidelity"]["score"] == 1.0
    assert keyed_user.get(f"/api/imports/{detail['id']}/source").status_code == 404
    assert keyed_user.post("/api/imports/text", json={"text": "too short"}).status_code == 422

    scan = upload(keyed_user, corpus.scanned_pdf(), "scan.pdf")
    assert scan.status_code == 202
    drain(app)
    detail = keyed_user.get(f"/api/imports/{scan.json()['id']}").json()
    assert detail["hasTextLayer"] is False
    # The fake reads only text layers, so it finds nothing here (the real model reads the page images). The
    # import still completes, with a placeholder role to replace, and no fidelity check is possible.
    assert detail["status"] == "ready" and detail["fidelity"]["available"] is False
    assert detail["document"]["work"][0]["company"] == "Your employer"
    assert any("No experience or projects" in f["message"] for f in detail["flags"])


@pytest.mark.parametrize(("data", "status", "code"), [
    (corpus.encrypted_pdf, 422, "encrypted"), (corpus.docm, 422, "macros"), (corpus.long_pdf, 422, "too_many_pages"),
    (lambda: b"plain text, not a file", 422, "unsupported"),
])
def test_bad_uploads_are_refused_up_front(app: FastAPI, keyed_user: Api, data: Any, status: int, code: str) -> None:
    response = upload(keyed_user, data())
    assert response.status_code == status and response.json()["code"] == code
    assert keyed_user.get("/api/imports").json() == []
    assert not list((app.state.settings.data_dir / "uploads").rglob("*.*"))


def test_upload_limits(app: FastAPI, keyed_user: Api) -> None:
    big = upload(keyed_user, b"%PDF-" + b"0" * (5 * 1024 * 1024))
    assert big.status_code == 413
    huge = upload(keyed_user, b"0" * (7 * 1024 * 1024))
    assert huge.status_code == 413 and huge.json()["detail"] == "The request is too large."


def test_imports_need_a_key_and_respect_the_cap(app: FastAPI, make_user: Any) -> None:
    user = make_user("nokey@example.com")
    response = upload(user, corpus.text_pdf())
    assert response.status_code == 409 and response.json()["code"] == "no_api_key"
    user.put("/api/account/api-key", json={"apiKey": "sk-ant-api03-another-test-key-0123456789AB",
                                           "password": USER_PASSWORD})
    user.put("/api/account/ai-settings", json={"monthlyCapUsd": 0.01})
    response = upload(user, corpus.text_pdf())
    assert response.status_code == 409 and response.json()["code"] == "cap_reached"


class FailingImportLLM(FakeLLM):
    fail = True

    async def complete(self, request: LLMRequest) -> LLMResult:
        if FailingImportLLM.fail:
            from app.ai.llm import Usage

            raise LLMError("truncated", "The response was cut off.", usage=Usage(1000, 32000), model=request.model)
        return await super().complete(request)


def test_failed_imports_record_usage_and_can_be_retried(app: FastAPI, keyed_user: Api) -> None:
    FailingImportLLM.fail = True
    app.state.llm_factory = FailingImportLLM
    response = upload(keyed_user, corpus.text_pdf())
    drain(app)
    import_id = response.json()["id"]
    detail = keyed_user.get(f"/api/imports/{import_id}").json()
    assert detail["status"] == "failed" and detail["error"]["code"] == "truncated" and detail["costUsd"] > 0
    stages = {s["stage"] for s in keyed_user.get("/api/account/usage").json()["byStage"]}
    assert "import" in stages
    assert keyed_user.put(f"/api/imports/{import_id}/review", json={"document": {}}).status_code == 409

    FailingImportLLM.fail = False
    assert keyed_user.post(f"/api/imports/{import_id}/retry").status_code == 202
    drain(app)
    assert keyed_user.get(f"/api/imports/{import_id}").json()["status"] == "ready"
    assert keyed_user.post(f"/api/imports/{import_id}/retry").status_code == 409


def test_progress_events(app: FastAPI, keyed_user: Api) -> None:
    response = upload(keyed_user, corpus.text_pdf())
    drain(app)
    events = keyed_user.get(f"/api/imports/{response.json()['id']}/events").text
    assert 'event: status\ndata: {"type": "status", "status": "ready"}' in events and "event: done" in events


def test_deleting_masters_and_accounts_removes_uploaded_files(app: FastAPI, keyed_user: Api) -> None:
    uploads = app.state.settings.data_dir / "uploads" / keyed_user.user["id"]
    detail = imported(app, keyed_user, corpus.text_pdf())
    master = keyed_user.post(f"/api/imports/{detail['id']}/accept",
                             json={"label": "m", "document": detail["document"]}).json()["master"]
    assert len(list(uploads.iterdir())) == 1
    assert keyed_user.delete(f"/api/master-resumes/{master['id']}").status_code == 204
    assert list(uploads.iterdir()) == []

    imported(app, keyed_user, corpus.full_docx(), "cv.docx")
    export = keyed_user.get("/api/account/export")
    names = zipfile.ZipFile(io.BytesIO(export.content)).namelist()
    assert any(n.startswith("uploads/") and n.endswith("cv.docx") for n in names)
    assert keyed_user.request("DELETE", "/api/account", json={"password": USER_PASSWORD}).status_code == 204
    assert not uploads.exists()


def test_interrupted_imports_are_marked_failed_on_restart(app: FastAPI, keyed_user: Api) -> None:
    from app.db.models import ImportJob
    from app.pipeline.importer import mark_interrupted

    response = upload(keyed_user, corpus.text_pdf())
    drain(app)
    import_id = response.json()["id"]

    async def interrupt() -> int:
        async with app.state.sessionmaker() as db:
            job = await db.get(ImportJob, import_id)
            assert job is not None
            job.status = "transcribing"
            await db.commit()
            return await mark_interrupted(db)

    assert app.state.test_portal.call(interrupt) == 1
    assert keyed_user.get(f"/api/imports/{import_id}").json()["error"]["code"] == "interrupted"


# ---- the Anthropic request for stage I ----------------------------------------------------------------------


@pytest.mark.anyio
async def test_pdfs_are_sent_as_document_blocks(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.ai import stages
    from tests.test_pipeline import _Message, _Stream

    client = AnthropicLLM("sk-ant-test-0000000000000000")
    seen: dict[str, Any] = {}
    draft = _draft().model_dump(by_alias=True)

    def fake_stream(**kwargs: Any) -> _Stream:
        seen.update(kwargs)
        import json

        return _Stream(_Message("end_turn", json.dumps(draft)))

    monkeypatch.setattr(client._client.beta.messages, "stream", fake_stream)
    result = await stages.transcribe(client, model="claude-opus-5-5", effort="high", kind="pdf", pdf=b"%PDF-1.7 x")
    assert result.value.basics.name == "Sam Lee" and result.prompt_version == "import@1"
    document, task = seen["messages"][0]["content"]
    assert document == {"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                        "data": "JVBERi0xLjcgeA=="}}
    assert task["text"].startswith('<task name="import">')
    assert seen["output_config"]["format"]["schema"]["properties"]["importNotes"]
    assert "Transcribe; don't improve" in seen["system"][0]["text"]

    seen.clear()
    await stages.transcribe(client, model="claude-opus-5-5", effort="high", kind="docx", text="Sam Lee\nEngineer")
    text = seen["messages"][0]["content"][0]["text"]
    assert text == '<resume_document format="docx">\nSam Lee\nEngineer\n</resume_document>'

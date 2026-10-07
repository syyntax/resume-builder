"""The import test corpus (DESIGN.md §17), generated at test time so no binary fixtures are committed.

Every document carries the same résumé, laid out the ways real ones are: a plain single-column PDF, a
two-column PDF with a sidebar, a LinkedIn "Save to PDF" export, a scan with no text layer, and a Word file
that keeps contact details in the header, skills in a table, a certification in a text box, and a link. The
hostile files (encrypted PDF, too many pages, macro-enabled .docm, ZIP bombs) must be rejected safely.
"""

from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fpdf import FPDF

NAME = "Jordan Rivera"
HEADLINE = "Senior Platform Engineer"
CONTACT = ["Denver, CO", "jordan.rivera@example.com", "(555) 010-0199", "linkedin.com/in/jrivera"]
SUMMARY = ("Platform engineer with 9 years of experience building reliable payment infrastructure on AWS and "
           "Kubernetes, with a focus on developer experience and cost.")
WORK = [
    ("Senior Platform Engineer", "Acme Payments", "Denver, CO", "Mar 2021", "Present", [
        "Led migration of 40 payment services from EC2 to EKS, cutting p99 latency 45% and infrastructure cost "
        "$380K per year.",
        "Built a canary release pipeline with Argo Rollouts that reduced failed releases by 60%.",
        "Mentored 6 engineers and ran the on-call rotation for the platform team.",
    ]),
    ("Software Engineer", "Initech", "Austin, TX", "Jun 2017", "Feb 2021", [
        "Designed a Terraform module library adopted by 12 product teams.",
        "Cut CI build times from 25 to 8 minutes by introducing remote caching.",
    ]),
    ("Junior Developer", "Globex", "Austin, TX", "2015", "2017", [
        "Maintained internal billing tools written in Python and PostgreSQL.",
    ]),
]
PROJECT = ("tf-guard", "Open-source Terraform policy linter", ["Reached 1,200 GitHub stars and 40 contributors."])
EDUCATION = "B.S. in Computer Science, University of Colorado Boulder, 2015"
CERTIFICATION = "Certified Kubernetes Administrator - CNCF (2023)"
SKILLS = [("Cloud & Infrastructure", "AWS, Kubernetes, Terraform"), ("Languages", "Python, Go, SQL")]

# Text that must survive the import word for word (the fidelity corpus checks these).
KEY_TEXT = [NAME, SUMMARY, *(b for *_, bullets in WORK for b in bullets), *PROJECT[2]]


def plain_text() -> str:
    lines = [NAME, HEADLINE, " | ".join(CONTACT), "", "SUMMARY", SUMMARY, "", "EXPERIENCE"]
    for position, company, location, start, end, bullets in WORK:
        lines += [f"{position}, {company}", f"{location} | {start} - {end}", *(f"- {b}" for b in bullets)]
    lines += ["", "PROJECTS", f"{PROJECT[0]} - {PROJECT[1]}", *(f"- {b}" for b in PROJECT[2]),
              "", "EDUCATION", EDUCATION, "", "CERTIFICATIONS", CERTIFICATION, "", "SKILLS",
              *(f"{category}: {items}" for category, items in SKILLS)]
    return "\n".join(lines)


# ---- PDFs ------------------------------------------------------------------------------------------------


def _pdf() -> FPDF:
    pdf = FPDF(format="letter")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    return pdf


def _write(pdf: FPDF, text: str, *, width: float = 0, size: float = 10, bold: bool = False, gap: float = 0) -> None:
    pdf.set_font("Helvetica", "B" if bold else "", size)
    pdf.multi_cell(width, size * 0.5, text, new_x="LMARGIN", new_y="NEXT")
    if gap:
        pdf.ln(gap)


def text_pdf() -> bytes:
    pdf = _pdf()
    for line in plain_text().splitlines():
        if not line:
            pdf.ln(2)
        elif line.isupper():
            _write(pdf, line, bold=True, size=11)
        else:
            _write(pdf, line)
    return bytes(pdf.output())


def two_column_pdf() -> bytes:
    """Main column first, then a sidebar: content-stream order is reading order, as Word and LaTeX write it."""
    pdf = _pdf()
    pdf.set_left_margin(75)
    pdf.set_xy(75, 15)
    _write(pdf, NAME, size=18, bold=True)
    _write(pdf, HEADLINE, size=12, gap=3)
    _write(pdf, "SUMMARY", bold=True, size=11)
    _write(pdf, SUMMARY, width=125, gap=3)
    _write(pdf, "EXPERIENCE", bold=True, size=11)
    for position, company, location, start, end, bullets in WORK:
        _write(pdf, f"{position}, {company}", width=125, bold=True)
        _write(pdf, f"{location} | {start} - {end}", width=125)
        for bullet in bullets:
            _write(pdf, f"- {bullet}", width=125)
        pdf.ln(2)
    pdf.set_left_margin(12)
    pdf.set_xy(12, 15)
    _write(pdf, "CONTACT", bold=True, size=11)
    for item in CONTACT:
        _write(pdf, item, width=58)
    pdf.ln(3)
    _write(pdf, "SKILLS", bold=True, size=11)
    for category, items in SKILLS:
        _write(pdf, f"{category}: {items}", width=58)
    pdf.ln(3)
    _write(pdf, "EDUCATION", bold=True, size=11)
    _write(pdf, EDUCATION, width=58, gap=3)
    _write(pdf, "CERTIFICATIONS", bold=True, size=11)
    _write(pdf, CERTIFICATION, width=58)
    return bytes(pdf.output())


def linkedin_pdf() -> bytes:
    """LinkedIn's "Save to PDF": sidebar of contact and top skills, then Company / Title / dates blocks."""
    pdf = _pdf()
    pdf.set_left_margin(12)
    pdf.set_xy(12, 15)
    _write(pdf, "Contact", bold=True, size=11)
    _write(pdf, "jordan.rivera@example.com", width=55)
    _write(pdf, "www.linkedin.com/in/jrivera (LinkedIn)", width=55, gap=3)
    _write(pdf, "Top Skills", bold=True, size=11)
    for skill in ("Kubernetes", "Terraform", "AWS"):
        _write(pdf, skill, width=55)
    pdf.ln(3)
    _write(pdf, "Certifications", bold=True, size=11)
    _write(pdf, "Certified Kubernetes Administrator", width=55)
    pdf.set_left_margin(75)
    pdf.set_xy(75, 15)
    _write(pdf, NAME, size=20, bold=True)
    _write(pdf, f"{HEADLINE} at Acme Payments", width=125)
    _write(pdf, "Denver, Colorado, United States", width=125, gap=4)
    _write(pdf, "Summary", bold=True, size=13)
    _write(pdf, SUMMARY, width=125, gap=4)
    _write(pdf, "Experience", bold=True, size=13)
    months = {"Mar": "March", "Jun": "June", "Feb": "February"}
    for position, company, _location, start, end, bullets in WORK:
        start_long = " ".join(months.get(part, part) for part in start.split())
        end_long = " ".join(months.get(part, part) for part in end.split())
        _write(pdf, company, width=125, bold=True)
        _write(pdf, position, width=125)
        _write(pdf, f"{start_long} - {end_long} (3 years)", width=125)
        for bullet in bullets:
            _write(pdf, bullet, width=125)
        pdf.ln(3)
    pdf.add_page()
    pdf.set_xy(75, 15)
    _write(pdf, "Education", bold=True, size=13)
    _write(pdf, "University of Colorado Boulder", width=125, bold=True)
    _write(pdf, "Bachelor of Science - BS, Computer Science · (2011 - 2015)", width=125)
    pdf.set_auto_page_break(auto=False)
    for page in (1, 2):
        pdf.page = page
        pdf.set_xy(95, 262)
        _write(pdf, f"Page {page} of 2", width=40)
    return bytes(pdf.output())


def scanned_pdf() -> bytes:
    """An image-only PDF, like a scan: no text layer at all."""
    from PIL import Image, ImageDraw

    image = Image.new("L", (1275, 1650), 255)
    draw = ImageDraw.Draw(image)
    y = 80
    for line in plain_text().splitlines():
        draw.text((90, y), line[:110], fill=0)
        y += 24
    buffer = io.BytesIO()
    image.save(buffer, format="PDF", resolution=150)
    return buffer.getvalue()


def encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(text_pdf())))
    writer.encrypt(user_password="secret", owner_password="owner")  # noqa: S106  (a test fixture)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def long_pdf(pages: int = 11) -> bytes:
    pdf = _pdf()
    for n in range(pages):
        if n:
            pdf.add_page()
        _write(pdf, f"Page {n + 1}: {SUMMARY}")
    return bytes(pdf.output())


# ---- Word -----------------------------------------------------------------------------------------------------


def _hyperlink(paragraph, url: str, text: str) -> None:  # noqa: ANN001
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}"><w:r><w:t>{text}</w:t></w:r></w:hyperlink>')
    paragraph._p.append(link)  # noqa: SLF001


def _text_box(paragraph, text: str) -> None:  # noqa: ANN001
    box = parse_xml(
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>'
        '<v:shape style="width:200pt;height:40pt"><v:textbox><w:txbxContent>'
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    )
    paragraph._p.append(box)  # noqa: SLF001


def full_docx() -> bytes:
    document = Document()
    header = document.sections[0].header
    header.paragraphs[0].text = NAME
    header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.add_paragraph(f"{HEADLINE} • {' • '.join(CONTACT)}")
    document.sections[0].footer.paragraphs[0].text = "Résumé of Jordan Rivera – references on request"

    document.add_heading("Summary", level=1)
    document.add_paragraph(SUMMARY)
    document.add_heading("Experience", level=1)
    for position, company, location, start, end, bullets in WORK:
        document.add_paragraph(f"{position}, {company}").runs[0].bold = True
        document.add_paragraph(f"{location} | {start} – {end}")
        for bullet in bullets:
            document.add_paragraph(bullet, style="List Bullet")
    document.add_heading("Projects", level=1)
    projects = document.add_paragraph(f"{PROJECT[0]} - {PROJECT[1]} ")
    _hyperlink(projects, "https://github.com/jrivera/tf-guard", "GitHub")
    for bullet in PROJECT[2]:
        document.add_paragraph(bullet, style="List Bullet")
    document.add_heading("Education", level=1)
    document.add_paragraph(EDUCATION)
    document.add_heading("Certifications", level=1)
    _text_box(document.add_paragraph(), CERTIFICATION)
    document.add_heading("Skills", level=1)
    table = document.add_table(rows=0, cols=2)
    for category, items in SKILLS:
        row = table.add_row().cells
        row[0].text, row[1].text = category, items
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _rezip(data: bytes, extra: dict[str, bytes], replace: dict[str, bytes] | None = None) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry.filename, (replace or {}).get(entry.filename, source.read(entry.filename)))
        for name, content in extra.items():
            target.writestr(name, content)
    return out.getvalue()


def docm() -> bytes:
    data = full_docx()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        types = archive.read("[Content_Types].xml").replace(
            b"wordprocessingml.document.main+xml", b"ms-word.document.macroEnabled.main+xml")
    return _rezip(data, {"word/vbaProject.bin": b"\x00" * 64}, {"[Content_Types].xml": types})


def zip_bomb(size: int = 60 * 1024 * 1024) -> bytes:
    """Declares its unpacked size honestly: rejected before anything is decompressed."""
    return _rezip(full_docx(), {"word/media/padding.xml": b"<a/>" + b" " * size})


def lying_zip_bomb(size: int = 60 * 1024 * 1024) -> bytes:
    """Claims a tiny unpacked size in both headers; the sandboxed worker must still refuse it."""
    data = bytearray(zip_bomb(size))
    name = b"word/media/padding.xml"
    for signature, offset in ((b"PK\x03\x04", 22), (b"PK\x01\x02", 24)):
        start = 0
        while (start := data.find(signature, start)) != -1:
            name_at = start + (30 if signature == b"PK\x03\x04" else 46)
            if data[name_at:name_at + len(name)] == name:
                data[start + offset:start + offset + 4] = (100).to_bytes(4, "little")
            start += 4
    return bytes(data)


def docx_with_dtd() -> bytes:
    data = full_docx()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        body = archive.read("word/document.xml")
    evil = body.replace(b"?>", b'?><!DOCTYPE w [<!ENTITY a "aaaaaaaaaa">]>', 1)
    return _rezip(data, {}, {"word/document.xml": evil})
